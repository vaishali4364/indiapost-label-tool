"""India Post Bulk Label Generator - Sharmex Global (v2)

Changes vs v1
  * Ship-To ("Post to ...") columns are read by their real eBay names and are ALWAYS used.
    Buyer address is only used as a flagged (red) last resort when Ship-To is completely empty.
  * Multi-buy fixed: pack size is read from "Variation details" (e.g. [Select:Pack of 5]),
    not only from the title; multi-line orders are merged into ONE parcel; same-address
    orders are flagged.
  * Product names are rebuilt as  Brand + Product + Variant + Size + Pack  (readable by anyone).
  * Every listing now gets a proper HS code / weight (category engine), instead of a generic
    "140 g / 33049990" fallback. Anything estimated is highlighted so you can verify it.
  * Optional overrides.xlsx lets you permanently correct any item (HS, weight, value, name).
  * Writes into your template.xlsx (keeps its formatting and the Information sheet) and gives
    a separate Review Report so the upload file stays clean.
"""
import base64
import io
import json
import math
import os
import re
import urllib.error
import urllib.request

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill

TEMPLATE_PATH = "template.xlsx"
OVERRIDES_PATH = "overrides.xlsx"  # optional, lives next to app.py
HISTORY_PATH = "order_history.csv"  # learned weights/prices, lives next to app.py
ESTIMATED_SOURCES = ("Estimated (category)", "UNKNOWN item")

SENDER = {
    "SENDER NAME": "NEELA SHARMA",
    "SENDER COMPANY": "SHARMEX GLOBAL",
    "SENDER ADD LINE 1": "1037, NEW MODEL TOWN",
    "SENDER ADD LINE 2": "PINJORE",
    "SENDER CITY": "PANCHKULA",
    "SENDER STATE": "HR",
    "SENDER COUNTRY NAME": "India",
    "SENDER COUNTRY CODE": "IN",
    "SENDER PINCODE": 134102,
    "SENDER EMAILID": "sharmexglobal@gmail.com",
    "SENDER MOBILE": "8629056095",
}

DEFAULT_RATES = {"GBP": 128.0, "USD": 96.0, "CAD": 70.0, "AUD": 62.0, "EUR": 110.0}

# destination text -> (code, India Post name, currency)
COUNTRIES = {
    "united kingdom": ("GB", "United Kingdom", "GBP"),
    "uk": ("GB", "United Kingdom", "GBP"),
    "great britain": ("GB", "United Kingdom", "GBP"),
    "gb": ("GB", "United Kingdom", "GBP"),
    "united states": ("US", "United States of America", "USD"),
    "united states of america": ("US", "United States of America", "USD"),
    "usa": ("US", "United States of America", "USD"),
    "us": ("US", "United States of America", "USD"),
    "canada": ("CA", "Canada", "CAD"),
    "australia": ("AU", "Australia", "AUD"),
    "germany": ("DE", "Germany", "EUR"),
    "ireland": ("IE", "Ireland", "EUR"),
    "france": ("FR", "France", "EUR"),
}

PINK = PatternFill("solid", start_color="FFB6C1", end_color="FFB6C1")  # multi-buy (weight)
YELLOW = PatternFill("solid", start_color="FFFF00", end_color="FFFF00")  # state = city
RED = PatternFill("solid", start_color="FF6B6B", end_color="FF6B6B")  # must fix
SKY = PatternFill("solid", start_color="87CEEB", end_color="87CEEB")  # county looked up online
ORANGE = PatternFill("solid", start_color="FFC000", end_color="FFC000")  # estimated item data


# ======================================================================
# helpers
# ======================================================================
def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def clean_str(v, fallback=""):
    if v is None:
        return fallback
    s = str(v).strip()
    if s.lower() in ("nan", "none", "nat"):
        return fallback
    return re.sub(r"\s+", " ", s)


def tidy_line(s):
    return re.sub(r"[\s,;]+$", "", re.sub(r"^[\s,;]+", "", clean_str(s)))


def money(v):
    s = re.sub(r"[^\d.]", "", clean_str(v))
    try:
        return float(s) if s else 0.0
    except ValueError:
        return 0.0


# ======================================================================
# 1. CSV loading
# ======================================================================
def read_orders_csv(raw_bytes):
    text = raw_bytes.decode("utf-8-sig", errors="ignore")
    lines = text.splitlines()
    idx = next(
        (i for i, l in enumerate(lines[:15]) if "order number" in l.lower()), -1
    )
    if idx == -1:
        raise ValueError("Could not find the 'Order number' header row in the CSV.")
    df = pd.read_csv(
        io.StringIO("\n".join(lines[idx:])),
        dtype=str,
        keep_default_na=False,
        on_bad_lines="skip",
    )
    df.columns = [str(c).strip() for c in df.columns]
    cols = {norm(c): c for c in df.columns}
    order_col = cols.get("ordernumber")
    title_col = cols.get("itemtitle")
    if not order_col or not title_col:
        raise ValueError("CSV needs 'Order number' and 'Item title' columns.")
    # multi-item orders: rows with a title but blank order number belong to the row above
    last = ""
    orders = []
    for o, t in zip(df[order_col], df[title_col]):
        o = o.strip()
        if o:
            last = o
        elif t.strip():
            o = last
        orders.append(o)
    df["_order"] = orders
    df = df[df["_order"].str.match(r"^\d{2}-\d{5}-\d{5}$")].copy()
    df = df.drop_duplicates()
    return df, cols


# ======================================================================
# 2. Ship-To resolution  (the part that used to break)
# ======================================================================
SHIP_PREFIXES = ["postto", "shipto", "shipping", "ship", "recipient", "delivery"]
FIELD_SUFFIX = {
    "name": ["name", "fullname"],
    "phone": ["phone", "mobile", "phonenumber", "contactphone"],
    "add1": ["address1", "addressline1", "street1", "add1"],
    "add2": ["address2", "addressline2", "street2", "add2"],
    "city": ["city", "town"],
    "state": ["state", "county", "stateprovince", "province", "region"],
    "zip": ["postcode", "zip", "zipcode", "postalcode"],
    "country": ["country", "countrycode"],
}


def candidates(prefixes, field):
    return [p + s for p in prefixes for s in FIELD_SUFFIX[field]]


def first_nonblank(group, colmap, keys):
    """colmap: norm(col)->col. Scan EVERY row of the order (not just the first)."""
    for k in keys:
        col = colmap.get(k)
        if not col:
            continue
        for v in group[col]:
            v = clean_str(v)
            if v:
                return v
    return ""


def resolve_address(group, colmap):
    fields = ["name", "phone", "add1", "add2", "city", "state", "zip", "country"]
    ship = {f: first_nonblank(group, colmap, candidates(SHIP_PREFIXES, f)) for f in fields}
    buyer = {f: first_nonblank(group, colmap, candidates(["buyer"], f)) for f in fields}
    ship_has_block = bool(ship["name"] or ship["add1"] or ship["zip"])
    if ship_has_block:
        addr, source = ship, "ship"
        # same-block rule: never mix buyer fields into a ship-to address,
        # except phone (buyer phone is only used if ship-to phone is missing)
        if not addr["phone"]:
            addr = {**addr, "phone": buyer["phone"]}
    else:
        addr, source = buyer, "buyer"
    for k in ("name", "add1", "add2", "city", "state"):
        addr[k] = tidy_line(addr[k])
    return addr, source, ship, buyer


_COUNTY_CACHE = {}


def lookup_county(zipcode, code, timeout=6):
    """Look up county/state from the postcode on the internet (free APIs, no key).
    GB: postcodes.io  -> admin_county, else admin_district, else region
    US: zippopotam.us -> state abbreviation
    Returns '' if offline / not found (caller then falls back to City)."""
    key = (zipcode, code)
    if key in _COUNTY_CACHE:
        return _COUNTY_CACHE[key]
    out = ""
    try:
        if code == "GB" and zipcode:
            url = "https://api.postcodes.io/postcodes/" + urllib.request.quote(zipcode.replace(" ", ""))
        elif code == "US" and re.fullmatch(r"\d{5}", zipcode or ""):
            url = f"https://api.zippopotam.us/us/{zipcode}"
        else:
            return ""
        with urllib.request.urlopen(url, timeout=timeout) as r:
            j = json.loads(r.read().decode("utf-8"))
        if code == "GB":
            res = j.get("result") or {}
            out = res.get("admin_county") or res.get("admin_district") or res.get("region") or ""
        else:
            out = (j.get("places") or [{}])[0].get("state abbreviation", "")
    except Exception:
        return ""  # do not cache failures
    _COUNTY_CACHE[key] = out
    return out


def norm_postcode(z, code):
    z = clean_str(z).upper()
    if code == "US":
        return z.split("-")[0].strip()
    if code == "GB":
        compact = re.sub(r"\s+", "", z)
        if len(compact) >= 5:
            return compact[:-3] + " " + compact[-3:]
    return z


# ======================================================================
# 3. Pack / size parsing
# ======================================================================
def variation_pack(var):
    """'[Select:Pack of 5]' / 'Pack=2' / 'Select=5 Black' / 'MPN:Pack of 10' -> int or None"""
    var = clean_str(var).strip("[] ")
    if not var:
        return None
    for part in re.split(r"[;,]", var):
        _, _, v = part.partition(":") if ":" in part else part.partition("=")
        v = v.strip()
        m = re.search(r"pack\s*of\s*(\d+)", v, re.I) or re.search(r"(\d+)\s*[- ]?pack\b", v, re.I)
        if m:
            return int(m.group(1))
        if re.fullmatch(r"\d{1,2}", v):
            return int(v)
        m = re.match(r"(\d{1,2})\s+(?!ml\b|g\b|gm\b|kg\b|oz\b|l\b)[a-z]", v, re.I)
        if m:
            return int(m.group(1))
    return None


def title_pack(title, allow_x=True):
    t = clean_str(title).lower()
    mult = 1
    m = re.match(r"^\s*(\d+)\s*x\s*(?=(?:pack|set|box)\s*of)", t)
    if m:
        mult = int(m.group(1))
    nums = []
    pats = [
        r"(?:packs?|set|box|bundle|combo|lot)\s*of\s*(\d+)",
        r"\((\d+)\s*total\)",
        r"\b(\d+)\s*[- ]?\s*(?:pack|pcs|pieces|pc|pk)\b",
    ]
    if allow_x:
        pats += [
            r"\bx\s*(\d+)\b(?!\s*(?:ml|g|gm|kg|gram|l\b))",
            r"(?<![\d.])(\d+)\s*x\s*(?=\d+(?:\.\d+)?\s*(?:ml|g|gm|kg))",
        ]
    for p in pats:
        nums += [int(x) for x in re.findall(p, t)]
    if "twin pack" in t:
        nums.append(2)
    for a, b in re.findall(r"(\d+)\s*packs?\s*of\s*(\d+)", t):
        nums.append(int(a) * int(b))
    m = re.match(r"^\s*(\d+)\s+(?=[a-z])(?!ml|g\b|gm|kg|in\b)", t)
    if m and 2 <= int(m.group(1)) <= 12:
        nums.append(int(m.group(1)))
    n = max(nums) if nums else 1
    return max(1, min(mult * n, 60))


def detect_pack(title, variation):
    v = variation_pack(variation)
    if v:
        return max(1, min(v, 60)), "variation"
    return title_pack(title), "title"


SIZE_RE = re.compile(r"(?<![\d.])(\d+(?:\.\d+)?)(?:\s*(kg|gms?|grams?|g|ml|ltr)|(l))(?![a-z'’])", re.I)


def parse_size(title):
    """returns (grams_or_ml_total, display_text) ; sums '340ml & Conditioner 180ml' sets"""
    t = clean_str(title)
    ms = list(SIZE_RE.finditer(t))
    if not ms:
        return 0.0, ""

    def val(m):
        v = float(m.group(1))
        return v * 1000 if (m.group(2) or "l").lower() in ("kg", "ltr", "l") else v

    def txt(m):
        n = m.group(1).rstrip("0").rstrip(".") if "." in m.group(1) else m.group(1)
        u = (m.group(2) or "l").lower()
        u = "g" if u.startswith("g") else u
        return f"{n}{u}"

    total, shown = val(ms[0]), txt(ms[0])
    if len(ms) > 1:
        between = t[ms[0].end():ms[1].start()]
        if re.fullmatch(r"\s*(?:&|\+|and)\s*(?:[A-Za-z ]{0,20})?\s*", between):
            total += val(ms[1])
            shown += "+" + txt(ms[1])
    return total, shown


# ======================================================================
# 4. Product specs: rules (your tuned numbers) -> categories -> unknown
# ======================================================================
def S(label, desc, hs, wt, val, conf="rule"):
    return {"label": label, "desc": desc, "hs": hs, "wt": wt, "val": val, "conf": conf}


def tuned_rules(t, pack):
    """Your original hand-tuned weights/values. t = lower-case title."""
    if "olivia" in t and "bleach" in t:
        return S("Bleach Cream", "Olivia Herbal Bleach Cream", "33049910", 49 * pack, 55 * pack)
    if "betadine" in t:
        return S("Antiseptic", "Betadine Povidone-iodine", "30049099",
                 248 if pack == 2 else 140 * pack, 90 * pack)
    if "bella vita" in t or "bellavita" in t:
        hs = "33030040"
        bp = title_pack(t, allow_x=False)
        if "intense" in t and ("ceo" in t or "100" in t):
            return S("EDP", "Bellavita CEO Intense 100ml", hs, 448 * bp, 450 * bp)
        if any(k in t for k in ["combo", "2 pack", "2 x 100ml", "ceo + goat", "ceo and goat"]):
            return S("EDP Combo", "Bellavita CEO & GOAT 2Pk", hs, 848, 700)
        if any(k in t for k in ["gift set", "4 x 20ml", "4x20ml", "all star"]):
            if bp == 2 or "pack of 2" in t or "2 bella vita" in t or "set 2" in t:
                return S("Gift Set", "Bellavita Gift Set 4x20ml 2Pk", hs, 848, 700)
            return S("Gift Set", "Bellavita Gift Set 4x20ml", hs, 398, 350)
        if any(k in t for k in ["2 x 20ml", "2x20ml", "white and honey", "white oud & honey",
                                "fresh + white", "date and senorita", "ceo + white", "date & glam"]):
            return S("Perfume Duo", "Bellavita Perfume Duo 2x20ml", hs, 198, 360)
        if "20ml" in t:
            return S("Perfume", "Bellavita Pocket Perfume 20ml", hs, 98 * bp, 180 * bp)
        if any(k in t for k in ["dark oud", "oud dark", "oud gold"]):
            return S("EDP", "Bellavita Dark Oud EDP 100ml", hs, 498 * bp, 486 * bp)
        return S("EDP", "Bellavita Eau De Parfum 100ml", hs, 398 * bp, 350 * bp)
    if any(k in t for k in ["detan", "moroccon", "face pack"]):
        return S("De-Tan Mask", "Sayy Moroccan DeTan Mask", "33049090", 140 * pack, 210 * pack)
    if "elvive" in t or "biolage" in t:
        if "combo" in t:
            return S("Hair Combo", "Loreal Elvive Hair Combo", "33059090", 548, 506)
        wt = {2: 240, 3: 348, 4: 490}.get(pack, 140 * pack) if pack > 1 else 140
        return S("Hair Serum", "Loreal Paris Elvive Serum", "33059090", wt, 434 * pack)
    if any(k in t for k in ["gluta hya", "gluta-hya", "vaseline"]) and "lip" not in t:
        if any(k in t for k in ["bundle", "pack of 3", "(3)"]) or pack == 3:
            return S("Body Lotion", "Vaseline Gluta Hya 3Pk", "33049090", 645, 826)
        if "400ml" in t or "sun protect" in t:
            return S("Body Lotion", "Vaseline Body Lotion 400ml", "33049090", 480 * pack, 340 * pack)
        if pack > 1:
            return S("Body Lotion", "Vaseline Gluta Hya Lotion", "33049090",
                     398 if pack == 2 else 249 * pack, 255 * pack)
        return S("Body Lotion", "Vaseline Gluta Hya Lotion", "33049090", 249, 255)
    if ("serum" not in t and any(k in t for k in ["roll on", "roll-on", "rexona"])) or (
        "chemist at play" in t and ("roll" in t or "deodorant" in t)
    ):
        return S("Roll-On Deodorant", "Deodorant Underarm Roll On", "33072000", 98 * pack, 291 * pack)
    if "supradyn" in t:
        return S("Multivitamin", "Supradyn Multivitamin 60s", "21069099", 98 * pack, 250 * pack)
    if "sponge" in t:
        return S("Shoe Shine Sponge", "Kiwi Shoe Shine Sponge", "34051000",
                 198 if pack == 2 else 98 * pack, 170 * pack)
    if any(k in t for k in ["kiwi", "cherry blossom", "shoe polish"]):
        if any(k in t for k in ["instant", "liquid"]):
            return S("Shoe Polish", "Kiwi Instant Shoe Polish", "34051000",
                     198 if pack == 2 else 98 * pack, 106 * pack)
        return S("Shoe Polish", "Kiwi Shoe Polish Wax Tin", "34051000",
                 248 if pack == 2 else 120 * pack, 106 * pack)
    if "copper" in t:
        if any(k in t for k in ["pot", "jug"]):
            return S("Copper Water Pot", "Pure Copper Water Pot", "74198090", 326, 945)
        if "ball" in t:
            wt = 49 if pack <= 2 else (98 if pack <= 4 else 198)
            return S("Copper Ball", "Ayurvedic Copper Ball", "74198090", wt,
                     162 * (pack // 2 if pack > 1 else 1))
    if any(k in t for k in ["guard", "mach 3", "fusion", "vector", "razor"]):
        return S("Shaving Razor", "Gillette Shaving Razor", "82121010",
                 198 if pack == 2 else 98 * pack, 260 * pack)
    if any(k in t for k in ["toothbrush", "oral b", "oral-b"]):
        return S("Toothbrush", "Oral Care Toothbrush Pack", "96032100",
                 198 if pack == 2 else 98 * pack, 213 * pack)
    if "toothpaste" in t or "tooth paste" in t or "colgate" in t or ("dabur" in t and "paste" in t):
        return S("Toothpaste", "Dental Toothpaste Tube", "33061020", 198 * pack, 210 * pack)
    return None


# (regex, label, HS, default_wt_g, weight_from_size, value_inr_per_unit)
# HS codes below are best-effort defaults - items matched here are highlighted ORANGE for you to verify.
CATEGORIES = [
    (r"\bmoov\b", "Pain Relief Cream", "30049011", 90, True, 220),
    (r"tata tea|black tea", "Black Tea", "09024020", 540, True, 250),
    (r"\bhing\b|asafoetida", "Asafoetida (Hing) Powder", "13019013", 80, True, 150),
    (r"dant manjan|tooth ?powder", "Tooth Powder", "33061090", 130, True, 150),
    (r"copper.*bottle|bottle.*copper", "Copper Water Bottle", "74198090", 350, False, 800),
    (r"bleach", "Bleach Cream", "33049910", 60, True, 100),
    (r"copper.*glass|drinking glass|tumbler", "Copper Drinking Glass", "74198090", 150, False, 400),
    (r"copper", "Copper Utensil", "74198090", 300, False, 500),
    (r"toner", "Face Toner", "33049990", 120, True, 300),
    (r"incense|agarbatti|dhoop", "Incense Sticks", "33074100", 100, False, 150),
    (r"air freshener|freshener|\bglade\b|godrej aer", "Air Freshener", "33074900", 60, False, 150),
    (r"acupressure|sujok|massage ring", "Acupressure Finger Massage Ring", "90191090", 25, False, 150),
    (r"lip balm", "Lip Balm", "33041000", 35, False, 150),
    (r"lipstick|lip gloss|lip liner", "Lipstick", "33041000", 40, False, 150),
    (r"mascara", "Mascara", "33042000", 60, False, 250),
    (r"kajal|kohl|eye ?liner|insta liner", "Kajal / Eye Liner", "33042000", 35, False, 150),
    (r"root concealer|retouch|root touch", "Root Touch-Up Spray", "33059090", 150, False, 300),
    (r"foundation", "Foundation", "33049990", 80, True, 300),
    (r"concealer", "Concealer", "33049990", 40, False, 250),
    (r"primer", "Face Primer", "33049990", 40, False, 250),
    (r"\bbb\+? cream|bb cream", "BB Cream", "33049990", 60, True, 250),
    (r"sunscreen|uv defender|sun protect|nivea sun|sun lotion|sun cream", "Sunscreen", "33049990", 90, True, 350),
    (r"face wash|facial wash|cleanser", "Face Wash", "33049990", 120, True, 200),
    (r"hair removal|wax strip|depilat|\bveet\b", "Hair Removal", "33079090", 120, True, 250),
    (r"shaving cream|shave cream|shave gel|shaving gel|after ?shave|lather", "Shaving Cream / Gel", "33071010", 110, True, 200),
    (r"deodorant|body spray|\bdeo\b", "Deodorant Spray", "33072000", 190, True, 291),
    (r"hair colou?r|hair dye|colou?r naturals|highlight kit|mehndi powder|henna powder|godrej nupur", "Hair Colour / Henna", "33059090", 130, True, 300),
    (r"henna|mehndi|mehendi", "Henna Mehndi Cone", "33049990", 55, False, 150),
    (r"shampoo", "Shampoo", "33051090", 200, True, 300),
    (r"conditioner", "Conditioner", "33059090", 150, True, 300),
    (r"hair mask", "Hair Mask", "33059090", 150, True, 300),
    (r"hair serum|serum", "Serum", "33059090", 120, True, 300),
    (r"hair oil|cool oil|bringha", "Hair Oil", "33059090", 150, True, 250),
    (r"soap|bathing bar|bath bar", "Bath Soap", "34011190", 135, True, 100),
    (r"talc|body powder|baby powder", "Body Powder", "33049990", 300, True, 250),
    (r"parfum|perfume|eau de|\bedt\b|\bedp\b|cologne", "EDP", "33030040", 400, False, 350),
    (r"body lotion|lotion", "Body Lotion", "33049990", 250, True, 280),
    (r"moisturi[sz]er|cream|gel", "Cream / Gel", "33049990", 110, True, 250),
]


def category_specs(t, pack, size):
    for pat, label, hs, dflt, from_size, val in CATEGORIES:
        if re.search(pat, t):
            if label == "EDP" and size:
                unit_wt = 150 + 2.5 * size
            elif from_size and size:
                unit_wt = size * 1.1 + 30
            else:
                unit_wt = dflt
            if label == "Black Tea" and size:
                val = size * 0.5
            return S(label, "", hs, round(max(unit_wt, 25)) * pack, round(val) * pack, "category")
    return S("Personal Care Item", "", "33049990", 140 * pack, 250 * pack, "unknown")


# ======================================================================
# 5. Readable product names
# ======================================================================
BRANDS = [
    (r"bella ?vita", "Bellavita"), (r"l['’`]?or[eé]al", "L'Oreal"), (r"maybelline( new york)?", "Maybelline"),
    (r"lakm[eé]", "Lakme"), (r"pond['’`]?s", "Ponds"), (r"nivea", "Nivea"), (r"garnier", "Garnier"),
    (r"old spice", "Old Spice"), (r"gillette", "Gillette"), (r"colgate", "Colgate"), (r"oral[- ]?b", "Oral-B"),
    (r"tata sampann", "Tata Sampann"), (r"tata tea", "Tata Tea"), (r"cherry blossom", "Cherry Blossom"),
    (r"kiwi", "Kiwi"), (r"moov", "Moov"), (r"betadine", "Betadine"), (r"supradyn", "Supradyn"),
    (r"vaseline", "Vaseline"), (r"veet", "Veet"), (r"renee", "Renee"), (r"blue heaven", "Blue Heaven"),
    (r"himalaya( herbals)?", "Himalaya"), (r"lotus", "Lotus"), (r"godrej", "Godrej"), (r"olay", "Olay"),
    (r"fixderma", "Fixderma"), (r"sayy", "Sayy"), (r"olivia", "Olivia"), (r"hamam", "Hamam"),
    (r"chandrika", "Chandrika"), (r"pears", "Pears"), (r"lux", "Lux"), (r"fem", "Fem"), (r"streax", "Streax"),
    (r"indulekha", "Indulekha"), (r"navratna", "Navratna"), (r"kaveri", "Kaveri"), (r"wild stone", "Wild Stone"),
    (r"johnson['’`]?s", "Johnson's"), (r"denver", "Denver"), (r"glade", "Glade"), (r"wishcare", "WishCare"),
    (r"the derma co", "The Derma Co"), (r"insight", "Insight"), (r"mars", "MARS"), (r"deewal", "Deewal"),
    (r"the body shop", "The Body Shop"), (r"the ordinary", "The Ordinary"), (r"urban gabru", "Urban Gabru"),
    (r"nirmalaya", "Nirmalaya"), (r"chemist at play", "Chemist At Play"), (r"rexona", "Rexona"),
    (r"dabur", "Dabur"), (r"elvive", "L'Oreal"), (r"biolage", "Biolage"),
]
KEEP_UPPER = {"BPA", "UV", "BB", "BB+", "EDP", "EDT", "CEO", "G.O.A.T", "GOAT", "AHA", "DIY", "SPF", "PA", "XL", "II", "III"}
FLUFF = {
    "luxury", "collection", "premium", "authentic", "genuine", "best", "new", "free", "shipping", "choose",
    "pack", "of", "each", "and", "the", "with", "for", "in", "a", "by", "from", "&", "+", "-", "–",
    "long", "lasting", "men", "women", "mens", "womens", "men's", "women's",
}
CUT_AT = re.compile(
    r"\s[|–—]\s|\s-\s|\bexp\b|\bchoose\b|\bmultiple\b|\blong[- ]lasting\b|\bfree shipping\b|"
    r"\bsmudge ?proof\b|\bwith\b|\bfor (?:men|women|all|oily|dry|hair|skin|sensitive)\b|"
    r"\bmade in\b|\b24 ?h\b|\bmatte finish\b|\bpack\b|\bset of\b",
    re.I,
)


def titlecase(word):
    if "-" in word and not any(c.isdigit() for c in word):
        return "-".join(titlecase(p) for p in word.split("-") if p)
    if re.fullmatch(r"(?:PA|SPF|UV)[\d+]*", word.upper()) and any(c in word for c in "+0123456789"):
        return word.upper()
    if any(c.isdigit() for c in word) or word.upper() in KEEP_UPPER:
        return word.upper() if word.upper() in KEEP_UPPER else word
    if "'" in word or "’" in word:
        return word[:1].upper() + word[1:].lower()
    return word.capitalize()


def friendly_name(title, label, pack, size_txt, max_len, cut_after_size=True):
    """Brand + product words + product type + size + pack, e.g.
    'Maybelline Colossal Kajal Black 0.35g' / 'Pears Original Glycerin Soap 125g (6-Pack)'"""
    t = clean_str(title)
    brand = ""
    for pat, canon in BRANDS:
        m = re.search(r"(?<![a-z])" + pat + r"(?![a-z])", t, re.I)
        if m:
            brand = canon
            t = t[:m.start()] + " " + t[m.end():]
            break
    t = re.sub(r"\([^)]*\)|\[[^\]]*\]", " ", t)  # (100 ml) (Pack of 2) ... re-added below
    t = re.sub(r"\bSPF\s*(\d+)\+*", r"SPF\1", t, flags=re.I)
    t = re.sub(r"\b\d+\s*x\s*pack\s*of\s*\d+\b|\b\d+\s*packs?\s*of\s*\d+\b", " ", t, flags=re.I)
    t = re.sub(r"\b(?:pack|set|box|bundle|combo|lot)\s*of\s*\d+\b", " ", t, flags=re.I)
    t = re.sub(r"\b\d+\s*[- ]?\s*(?:pack|pcs|pk)\b|\bx\s*\d+\b|\b\d+\s*x\s*(?=\d)|\b\d+\s*x\b|\btwin pack\b", " ", t, flags=re.I)
    # everything after the first size is usually marketing copy
    m = SIZE_RE.search(t)
    if cut_after_size and m and len(t[:m.start()].split()) >= 2:
        t = t[:m.start()]
    t = SIZE_RE.sub(" ", t)
    t = re.sub(r"\beau de (?:parfum|perfume)\b", "EDP", t, flags=re.I)
    t = re.sub(r"(?<![\w%])100%", " ", t)
    m = CUT_AT.search(t)
    if m and m.start() > 6:
        t = t[:m.start()]
    m = re.match(r"^\s*(\d{1,2})\s+(?=[A-Za-z])", t)  # leading '3 Glade ...' count
    if m and int(m.group(1)) == pack:
        t = t[m.end():]
    raw = [w.strip(".,;:-") for w in re.split(r"[\s,/]+", t)]
    words, seen = [], set()
    for w in raw:
        lw = w.lower()
        if not w or re.fullmatch(r"[-–—|+&.:]+", w) or lw in FLUFF or (len(lw) >= 5 and lw in seen):
            continue
        seen.add(lw)
        words.append(titlecase(w))

    suffix = (" " + size_txt if size_txt else "") + (f" ({pack}-Pack)" if pack > 1 else "")
    head = brand + " " if brand else ""
    budget = max_len - len(suffix) - len(head)

    lab_words = [w[:4] for w in re.split(r"[\s/()-]+", label.lower()) if len(w) > 2 and w not in ("care", "item", "and")]
    if label == "EDP":
        lab_words = ["edp", "parfum", "perfume"]
    has_label = lambda ws: any(lw in " ".join(ws).lower() for lw in lab_words)

    def fit(ws, room):
        ws = list(ws)
        while len(" ".join(ws)) > room and len(ws) > 1:
            # drop the last word that is NOT the product-type word
            idx = next((i for i in range(len(ws) - 1, -1, -1) if not has_label([ws[i]])), len(ws) - 1)
            ws.pop(idx)
        return ws

    core = fit(words, budget)
    if label and not has_label(words[:len(core)]):
        core = fit(words, budget - len(label) - 1)
        if not has_label(core):
            core.append(label.split(" / ")[0])
    out = (head + " ".join(core)).strip() + suffix
    return re.sub(r"\s+", " ", out)[:max_len].strip() or clean_str(title)[:max_len]


# ======================================================================
# 6. Overrides (optional, permanent corrections)
# ======================================================================
def load_overrides(file_or_path):
    try:
        df = pd.read_excel(file_or_path, dtype=str).fillna("")
    except Exception:
        return {}
    df.columns = [norm(c) for c in df.columns]
    out = {}
    for _, r in df.iterrows():
        iid = re.sub(r"\.0$", "", r.get("itemid", "").strip())
        if not iid:
            continue
        out[(iid, r.get("pack", "").strip())] = {
            "name": r.get("name", "").strip(),
            "hs": r.get("hscode", "").strip(),
            "wt": money(r.get("weightg", "")),
            "val": money(r.get("valueinr", "")),
        }
    return out


def resolve_item(title, variation, item_id, cfg, overrides, learned=None):
    is_bv = bool(re.search(r"bella ?vita", clean_str(title), re.I))
    v = variation_pack(variation)
    if v:
        pack, pack_src = max(1, min(v, 60)), "variation"
    else:
        pack, pack_src = title_pack(title, allow_x=not is_bv), "title"
    size, size_txt = parse_size(title)
    mset = re.search(r"(\d)\s*x\s*(\d+)\s*ml", clean_str(title), re.I) if is_bv else None
    if mset:
        size_txt = f"{mset.group(1)}x{mset.group(2)}ml"
    spec = tuned_rules(clean_str(title).lower(), pack)
    if spec is None:
        spec = category_specs(clean_str(title).lower(), pack, size)
    name = friendly_name(title, spec["label"], 1 if mset else pack, size_txt, cfg["max_name"],
                         cut_after_size=not is_bv)
    source = {"rule": "Your tuned rule", "category": "Estimated (category)", "unknown": "UNKNOWN item"}[spec["conf"]]
    hs, wt, val = spec["hs"], spec["wt"], spec["val"]
    ov = overrides.get((item_id, str(pack))) or overrides.get((item_id, ""))
    if ov:
        scale = 1 if (item_id, str(pack)) in overrides else pack
        hs = ov["hs"] or hs
        wt = ov["wt"] * scale if ov["wt"] else wt
        val = ov["val"] * scale if ov["val"] else val
        name = ov["name"] or name
        source = "Override file"
    elif learned:
        L = learned["exact"].get((item_id, pack))
        P = learned["per_pack"].get(item_id)
        if L:
            hs, wt, val, name = L["hs"] or hs, L["wt"], L["val"], L["name"] or name
            source = "Learned (history)"
        elif P:
            hs, wt, val = P["hs"] or hs, P["wt"] * pack, P["val"] * pack
            source = "Learned (history, per-pack)"
    return {
        "name": name, "hs": hs, "wt": int(round(wt)), "val": int(round(val)),
        "pack": pack, "pack_src": pack_src, "source": source,
        "estimated": source in ESTIMATED_SOURCES,
    }


# ======================================================================
# 6b. HSTUS code (US orders) and learned history
# ======================================================================
def hstus_code(hs, dest_code):
    """US orders only: HSN/HS code + '00' (8 digits -> 10 digits)."""
    if dest_code != "US":
        return ""
    d = re.sub(r"\D", "", clean_str(hs))
    if not d:
        return ""
    return d[:10] if len(d) >= 10 else d + "00"


HISTORY_COLS = ["saved_at", "order", "destination", "item_id", "pack", "qty", "title", "name", "hs", "hstus",
                "unit_wt", "unit_inr", "line_wt", "line_inr", "sold_unit", "currency", "fob_total", "source",
                "confirmed"]


def load_history(src):
    try:
        df = pd.read_csv(src, dtype=str).fillna("")
    except Exception:
        return pd.DataFrame(columns=HISTORY_COLS)
    for c in HISTORY_COLS:
        if c not in df.columns:
            df[c] = ""
    return df[HISTORY_COLS]


def merge_history(old, new):
    both = pd.concat([old, new], ignore_index=True)
    return both.drop_duplicates(["order", "item_id", "pack"], keep="last").reset_index(drop=True)


def learn_from_history(hist):
    """Only rows you confirmed/edited are trusted. Exact (item, pack) match first; otherwise per-pack rate x pack."""
    exact, per_pack = {}, {}
    if hist is None or hist.empty:
        return {"exact": exact, "per_pack": per_pack}
    h = hist[hist["confirmed"].astype(str).str.lower() == "true"].sort_values("saved_at", kind="stable")
    for _, r in h.iterrows():
        try:
            pack = max(1, int(float(r["pack"] or 1)))
            wt, val = float(r["unit_wt"]), float(r["unit_inr"])
        except ValueError:
            continue
        if wt <= 0:
            continue
        exact[(r["item_id"], pack)] = {"name": r["name"], "hs": r["hs"], "wt": wt, "val": val}
        per_pack[r["item_id"]] = {"hs": r["hs"], "wt": wt / pack, "val": val / pack}
    return {"exact": exact, "per_pack": per_pack}


def history_from_lines(lines_df):
    if lines_df is None or lines_df.empty:
        return pd.DataFrame(columns=HISTORY_COLS)
    now = pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")
    return pd.DataFrame({
        "saved_at": now, "order": lines_df["Order"], "destination": lines_df["Destination"],
        "item_id": lines_df["Item number"], "pack": lines_df["Pack"], "qty": lines_df["Qty"],
        "title": lines_df["eBay title"], "name": lines_df["India Post name"], "hs": lines_df["HS"],
        "hstus": lines_df["HSTUS"], "unit_wt": lines_df["Unit wt (g)"], "unit_inr": lines_df["Unit INR"],
        "line_wt": lines_df["Line wt (g)"], "line_inr": lines_df["Line INR"], "sold_unit": lines_df["Sold unit"],
        "currency": lines_df["Currency"], "fob_total": lines_df["FOB total"], "source": lines_df["Source"],
        "confirmed": lines_df["Confirmed"].astype(bool).astype(str),
    }).astype(str)[HISTORY_COLS]


def github_enabled():
    try:
        return bool(st.secrets.get("GITHUB_TOKEN")) and bool(st.secrets.get("GITHUB_REPO"))
    except Exception:
        return False


def github_push(path, data, message):
    """Optional: commit a file to your repo (needs GITHUB_TOKEN + GITHUB_REPO in Streamlit secrets)."""
    tok, repo = st.secrets["GITHUB_TOKEN"], st.secrets["GITHUB_REPO"]
    branch = st.secrets.get("GITHUB_BRANCH", "main")
    url = f"https://api.github.com/repos/{repo}/contents/{path}"
    hdr = {"Authorization": f"Bearer {tok}", "Accept": "application/vnd.github+json", "User-Agent": "label-tool"}
    sha = None
    try:
        with urllib.request.urlopen(urllib.request.Request(f"{url}?ref={branch}", headers=hdr), timeout=15) as r:
            sha = json.loads(r.read().decode())["sha"]
    except urllib.error.HTTPError as e:
        if e.code != 404:
            raise
    body = {"message": message, "content": base64.b64encode(data).decode(), "branch": branch}
    if sha:
        body["sha"] = sha
    req = urllib.request.Request(url, data=json.dumps(body).encode(), headers=hdr, method="PUT")
    with urllib.request.urlopen(req, timeout=20):
        pass


# ======================================================================
# 7. Order -> parcel
# ======================================================================
def currency_from_symbol(s, fallback):
    s = clean_str(s)
    if "£" in s:
        return "GBP"
    if "€" in s:
        return "EUR"
    if "$" in s:
        return {"CAD": "CAD", "AUD": "AUD"}.get(fallback, "USD")
    return fallback


def process(df, cols, cfg, overrides, learned=None):
    today = pd.Timestamp.now().strftime("%Y-%m-%d")
    col = lambda n: cols.get(norm(n))
    c_item, c_title, c_qty = col("Item number"), col("Item title"), col("Quantity")
    c_sold, c_var = col("Sold for"), col("Variation details")
    c_disp = col("Dispatched on date")

    parcels, lines_out, sub_rows = [], [], []
    marks = {}  # (serial, ArticleDetails column) -> fill
    sub_marks = []  # (subpiece row index, fill)
    addr_index = {}

    colmap = {norm(c): c for c in df.columns}
    serial = 0
    for order_id, grp in df.groupby("_order", sort=False):
        serial += 1
        addr, source, ship, buyer = resolve_address(grp, colmap)
        flags = []  # (level, text)

        # --- destination
        craw = addr["country"] or buyer["country"]
        info = COUNTRIES.get(craw.lower())
        if info:
            code, cname, cur_default = info
        else:
            code, cname, cur_default = "", craw or "?", "GBP"
            flags.append(("error", f"Country '{craw}' not recognised - set code/rate manually"))
        zipcode = norm_postcode(addr["zip"], code)

        if source == "buyer":
            flags.append(("error", "Ship-To address missing in CSV - BUYER address used, please check"))
            for f in ("RECEIVER NAME", "RECEIVER ADD LINE 1", "RECEIVER ZIPCODE"):
                marks[(serial, f)] = RED
        for key, colname, label in [("name", "RECEIVER NAME", "name"), ("add1", "RECEIVER ADD LINE 1", "address line 1"),
                                    ("city", "RECEIVER CITY", "city"), ("zip", "RECEIVER ZIPCODE", "postcode")]:
            if not (zipcode if key == "zip" else addr[key]):
                flags.append(("error", f"Receiver {label} is blank"))
                marks[(serial, colname)] = RED
        if code == "GB" and zipcode and not re.fullmatch(r"[A-Z]{1,2}\d[A-Z\d]? \d[A-Z]{2}", zipcode):
            flags.append(("warn", f"UK postcode '{zipcode}' looks wrong"))
            marks[(serial, "RECEIVER ZIPCODE")] = RED
        if code == "US" and zipcode and not re.fullmatch(r"\d{5}", zipcode):
            flags.append(("warn", f"US ZIP '{zipcode}' looks wrong"))
            marks[(serial, "RECEIVER ZIPCODE")] = RED

        state = addr["state"]
        if not state:
            found = lookup_county(zipcode, code) if cfg.get("lookup") else ""
            if found:
                state = found
                marks[(serial, "RECEIVER STATE")] = SKY
                flags.append(("info", f"County blank - '{found}' looked up online from postcode {zipcode} (sky blue) - please verify"))
            else:
                state = addr["city"]
                marks[(serial, "RECEIVER STATE")] = YELLOW
                flags.append(("info", "County blank and lookup unavailable - City used (yellow)"))

        phone = re.sub(r"\D", "", addr["phone"])
        if len(phone) < 8:
            flags.append(("warn", "Receiver phone missing - placeholder 0000000000 used"))
            phone = "0000000000"
            marks[(serial, "RECEIVER MOBILE NO")] = RED

        if source == "ship" and buyer["add1"] and (
            norm(buyer["add1"]) != norm(addr["add1"]) or norm(buyer["zip"]) != norm(addr["zip"])
        ):
            flags.append(("info", "Ship-To differs from buyer's own address (correct - Ship-To used)"))

        if c_disp and any(clean_str(v) for v in grp[c_disp]):
            flags.append(("info", "Already marked dispatched on eBay"))

        # --- lines
        rate_cur = cur_default
        parcel_wt = parcel_val = 0
        multi = len(grp) > 1
        order_lines = []
        for lsn, (_, row) in enumerate(grp.iterrows(), start=1):
            title = clean_str(row[c_title])
            if not title:
                continue
            item_id = re.sub(r"\.0$", "", clean_str(row[c_item]))
            try:
                qty = max(1, int(float(clean_str(row[c_qty]) or 1)))
            except ValueError:
                qty = 1
            variation = clean_str(row[c_var]) if c_var else ""
            it = resolve_item(title, variation, item_id, cfg, overrides, learned)
            cur = currency_from_symbol(row[c_sold] if c_sold else "", cur_default)
            rate_cur = cur
            rate = cfg["rates"].get(cur, DEFAULT_RATES.get(cur, 100.0))
            unit_fob = money(row[c_sold]) if c_sold else 0.0
            if not unit_fob:
                unit_fob = round(it["val"] / qty / rate, 2)
            tot_wt, tot_val = it["wt"] * qty, it["val"] * qty
            parcel_wt += tot_wt
            parcel_val += tot_val
            if qty > 1 or it["pack"] > 1:
                multi = True
            sub_rows.append({
                "SERIAL NUMBER REF ": serial, "HS CODE": it["hs"], "CTH CODE": it["hs"],
                "HS DESCRIPTION": it["name"], "SP UNIT CD": "PC", "SP COUNT": qty,
                "SP WEIGHT TOTAL": tot_wt, "SP NET WEIGHT": tot_wt,
                "SP ORIGIN COUNTRY CODE": "IN", "SP ORIGIN CURRENCY CODE": "INR",
                "SP COMM INVOICE NO": order_id, "SP COM INVOICE DATE(DD-MM-YYYY)": today,
                "SP INVOICE LSN": lsn, "SP INV CURRENCY CODE": cur,
                "SP INV EXCHANGE RATE": int(rate), "SP ASBL FOB VALUE": round(unit_fob * qty, 2),
                "SP ASBL INR VAL": tot_val, "SP TAX INVOICE NO": order_id, "SP TAX INVOICE DATE": today,
                "SP INV VALUE PU": it["val"], "SP INV VALUE TOTAL": tot_val,
                "ecommerce_url": "Ebay.com", "ecommerce_sku": item_id,
                "HSTUS CODE": hstus_code(it["hs"], code),
            })
            if it["estimated"]:
                sub_marks.append((len(sub_rows) - 1, ORANGE))
            order_lines.append({
                "Serial": serial, "Order": order_id, "Item number": item_id, "eBay title": title,
                "Variation": variation, "Qty": qty, "Pack": it["pack"], "Pack from": it["pack_src"],
                "India Post name": it["name"], "HS": it["hs"], "Unit wt (g)": it["wt"], "Line wt (g)": tot_wt,
                "Line INR": tot_val, "Source": it["source"],
                "Unit INR": it["val"], "Destination": code, "HSTUS": hstus_code(it["hs"], code),
                "Sold unit": unit_fob, "Currency": cur, "FOB total": round(unit_fob * qty, 2),
            })
            if it["source"] == "UNKNOWN item":
                flags.append(("warn", f"Item {item_id}: product type not recognised - check HS/weight"))
            elif it["estimated"]:
                flags.append(("info", f"Item {item_id}: weight/HS estimated from category (orange)"))
        lines_out += order_lines
        if not order_lines:
            serial -= 1
            continue
        if multi:
            marks[(serial, "PHYSICAL WEIGHT")] = PINK
            flags.append(("info", "Multi-buy / multi-pack parcel (pink weight)"))

        key = (norm(addr["name"]), norm(addr["add1"]), norm(zipcode))
        addr_index.setdefault(key, []).append((serial, order_id))

        parcels.append({
            "SERIAL NUMBER": serial, "DESTINATION COUNTRY NAME": cname, "DESTINATION COUNTRY CODE": code,
            "MAIL NATURE TYPE": 11, "MAIL TRANSPORT TYPE": "AMS", "PHYSICAL WEIGHT": parcel_wt,
            "DECLARED VALUE": parcel_val, "NON DELIVERY INSTRUCTIONS": "P", "INCOTERMS": "DDP",
            **SENDER,
            "RECEIVER NAME": addr["name"], "RECEIVER ADD LINE 1": addr["add1"],
            "RECEIVER ADD LINE 2": addr["add2"], "RECEIVER CITY": addr["city"], "RECEIVER STATE": state,
            "RECEIVER ZIPCODE": zipcode, "RECEIVER MOBILE NO": phone, "BULK REF": order_id,
            "DROP OFF PINCODE": 136118, "PBE TYPE": "PBEIII", "PBE FILING": "SELF",
            "_flags": flags, "_ship_source": source, "_buyer": buyer, "_currency": rate_cur,
            "_multi": multi,
        })

    # same address across different orders
    for key, lst in addr_index.items():
        if len(lst) > 1 and key[1]:
            for s, o in lst:
                others = ", ".join(x[1] for x in lst if x[0] != s)
                parcels[s - 1]["_flags"].append(("warn", f"Same address as other order(s): {others} (could be combined?)"))
    return parcels, sub_rows, sub_marks, marks, lines_out


# ======================================================================
# 8. Excel output
# ======================================================================
# columns India Post's newer templates may call slightly differently
ALIAS_KEYS = {"incoterm": "INCOTERMS", "hstus": "HSTUS CODE"}
REQUIRED_COLS = {"ArticleDetails": ["INCOTERMS"], "SubPieces": ["HSTUS CODE"]}


def build_workbook(parcels, sub_rows, sub_marks, marks):
    """Returns (xlsx bytes, [columns that were missing in template.xlsx and were added at the end])."""
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    added = []
    for name, rows in (("ArticleDetails", parcels), ("SubPieces", sub_rows)):
        ws = wb[name]
        header = {norm(c.value): c.column for c in ws[1] if c.value}
        for extra in REQUIRED_COLS[name]:
            n = norm(extra)
            have = n in header or any(ALIAS_KEYS.get(h) == extra for h in header)
            if not have:
                last = max(header.values())
                cell = ws.cell(1, last + 1, extra)
                prev = ws.cell(1, last)
                cell._style = prev._style
                header[n] = last + 1
                added.append(f"{name}: {extra}")
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        for r_i, row in enumerate(rows, start=2):
            rowmap = {norm(k): k for k in row}
            for hn, ci in header.items():
                key = rowmap.get(hn) or (ALIAS_KEYS.get(hn) if ALIAS_KEYS.get(hn) in row else None)
                v = row.get(key) if key else None
                if v == "":
                    v = None
                ws.cell(r_i, ci, v)
        if name == "ArticleDetails":
            for (serial, colname), fill in marks.items():
                if norm(colname) in header:
                    ws.cell(serial + 1, header[norm(colname)]).fill = fill
        else:
            for idx, fill in sub_marks:
                ws.cell(idx + 2, header[norm("HS DESCRIPTION")]).fill = fill
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), added


def apply_edits(parcels, sub_rows, lines_df, edited):
    """Apply your edits (name, HS, unit weight, unit INR) to the SubPieces + parcel totals.
    Returns (parcels, sub_rows, sub_marks, final_lines)."""
    new_rows = [dict(r) for r in sub_rows]
    final = lines_df.copy()
    final["Confirmed"] = False
    sub_marks = []
    for i in range(len(new_rows)):
        o, e = lines_df.iloc[i], edited.iloc[i]
        try:
            wt, val = float(e["Unit wt (g)"]), float(e["Unit INR"])
        except (TypeError, ValueError):
            wt, val = float(o["Unit wt (g)"]), float(o["Unit INR"])
        hs, name = clean_str(e["HS"]), clean_str(e["India Post name"])
        changed = (abs(wt - float(o["Unit wt (g)"])) > 1e-9 or abs(val - float(o["Unit INR"])) > 1e-9
                   or hs != clean_str(o["HS"]) or name != clean_str(o["India Post name"]))
        conf = bool(e["Confirmed"]) or changed
        r = new_rows[i]
        qty = r["SP COUNT"]
        tw, tv = int(round(wt * qty)), int(round(val * qty))
        r.update({"HS CODE": hs, "CTH CODE": hs, "HS DESCRIPTION": name, "SP WEIGHT TOTAL": tw, "SP NET WEIGHT": tw,
                  "SP ASBL INR VAL": tv, "SP INV VALUE PU": int(round(val)), "SP INV VALUE TOTAL": tv,
                  "HSTUS CODE": hstus_code(hs, o["Destination"])})
        if o["Source"] in ESTIMATED_SOURCES and not conf:
            sub_marks.append((i, ORANGE))
        final.iloc[i, final.columns.get_loc("India Post name")] = name
        for col, v in (("HS", hs), ("Unit wt (g)", int(round(wt))), ("Unit INR", int(round(val))),
                       ("Line wt (g)", tw), ("Line INR", tv), ("HSTUS", hstus_code(hs, o["Destination"]))):
            final.iloc[i, final.columns.get_loc(col)] = v
        final.iloc[i, final.columns.get_loc("Confirmed")] = conf or str(o["Source"]).startswith(("Learned", "Override"))
    wt_by, val_by = {}, {}
    for r in new_rows:
        k = r["SERIAL NUMBER REF "]
        wt_by[k] = wt_by.get(k, 0) + r["SP WEIGHT TOTAL"]
        val_by[k] = val_by.get(k, 0) + r["SP ASBL INR VAL"]
    new_parcels = []
    for p in parcels:
        p2 = dict(p)
        p2["PHYSICAL WEIGHT"] = wt_by.get(p["SERIAL NUMBER"], p["PHYSICAL WEIGHT"])
        p2["DECLARED VALUE"] = val_by.get(p["SERIAL NUMBER"], p["DECLARED VALUE"])
        new_parcels.append(p2)
    return new_parcels, new_rows, sub_marks, final


def build_report(parcels, lines_out):
    rows = []
    for p in parcels:
        lv = {"error": 0, "warn": 1, "info": 2}
        status = "CHECK" if any(l == "error" for l, _ in p["_flags"]) else (
            "Review" if any(l == "warn" for l, _ in p["_flags"]) else "OK")
        b = p["_buyer"]
        rows.append({
            "Serial": p["SERIAL NUMBER"], "Order": p["BULK REF"], "Status": status,
            "Ship-To name": p["RECEIVER NAME"], "Address 1": p["RECEIVER ADD LINE 1"],
            "Address 2": p["RECEIVER ADD LINE 2"], "City": p["RECEIVER CITY"], "State": p["RECEIVER STATE"],
            "Postcode": p["RECEIVER ZIPCODE"], "Country": p["DESTINATION COUNTRY CODE"],
            "Phone": p["RECEIVER MOBILE NO"], "Weight (g)": p["PHYSICAL WEIGHT"],
            "Declared INR": p["DECLARED VALUE"], "Multi-buy": "YES" if p["_multi"] else "",
            "Address source": "Ship-To" if p["_ship_source"] == "ship" else "BUYER (fallback!)",
            "Buyer address (info)": ", ".join(x for x in [b["add1"], b["city"], b["zip"]] if x),
            "Notes": " | ".join(t for l, t in sorted(p["_flags"], key=lambda x: lv[x[0]])),
            "_order_key": lv[min((l for l, _ in p["_flags"]), key=lambda x: lv[x], default="info")]
            if p["_flags"] else 3,
        })
    orders_df = pd.DataFrame(rows)
    return orders_df, pd.DataFrame(lines_out)


def overrides_template(lines_df):
    d = lines_df[lines_df["Source"].isin(["Estimated (category)", "UNKNOWN item"])]
    d = d.drop_duplicates(["Item number", "Pack"])
    return pd.DataFrame({
        "item_id": d["Item number"], "pack": d["Pack"], "title (info only)": d["eBay title"],
        "name": d["India Post name"], "hs_code": d["HS"],
        "weight_g": d["Line wt (g)"] / d["Qty"], "value_inr": d["Line INR"] / d["Qty"],
    })


# ======================================================================
# 9. UI
# ======================================================================
LEGEND = [
    ("#FFB6C1", "Pink", "Multi-buy / multi-pack: weight cell. Check the total weight."),
    ("#87CEEB", "Sky blue", "County/state was blank, looked up online. Please verify."),
    ("#FFFF00", "Yellow", "No county found, so the city was used."),
    ("#FF6B6B", "Red", "Must fix before uploading (missing/invalid address, phone, or buyer address used)."),
    ("#FFC000", "Orange", "Item name cell: weight/HS code is estimated. Please verify."),
]


def show_legend():
    rows = "".join(
        f'<div style="display:flex;align-items:center;gap:10px;margin:4px 0">'
        f'<span style="background:{c};color:#000;border:1px solid #888;border-radius:4px;'
        f'padding:2px 10px;min-width:78px;text-align:center;font-weight:600">{n}</span>'
        f'<span>{d}</span></div>'
        for c, n, d in LEGEND
    )
    with st.expander("🎨 Colour key (what the colours in the Excel mean)", expanded=True):
        st.markdown(rows, unsafe_allow_html=True)


# ======================================================================
# 10. eBay price calculator (same results as the reference calculator)
# ======================================================================
CALC_MARKETS = {
    "UK": dict(cur="GBP", sym="£", s0=236.0, ds=29.5, rate=115.0),
    "USA": dict(cur="USD", sym="$", s0=472.0, ds=41.3, rate=85.0),
    "Australia": dict(cur="AUD", sym="A$", s0=466.1, ds=53.1, rate=60.0),
    "Canada": dict(cur="CAD", sym="C$", s0=401.2, ds=47.2, rate=59.0),
}


def round_half_up(x, nd=2):
    f = 10 ** nd
    return math.floor(x * f + 0.5 + 1e-7) / f


def calc_shipping(weight_g, m):
    """Weight is rounded UP to the next 50 g slab; cost = first slab + step per extra slab (INR, GST incl.)."""
    if weight_g <= 0:
        return 0.0, 0
    n = math.ceil(weight_g / 50)
    return round_half_up(m["s0"] + m["ds"] * (n - 1)), n * 50


def calc_price99(price):
    """Next price ending in .99 that is above the 2-decimal price (11.81 -> 11.99, 11.99 -> 12.99)."""
    if price <= 0:
        return 0.0
    p = math.floor(price) + 0.99
    if p <= price + 1e-9:
        p += 1
    return round_half_up(p)


def calc_listing_price(cost, profit, weight_g, fee_pct, rate, m, ship_override=None):
    ship, slab = calc_shipping(weight_g, m)
    if ship_override:
        ship = ship_override
    base = cost + profit + ship
    price = round_half_up(base / (1 - fee_pct / 100) / rate)
    final = calc_price99(price)
    sale_inr = final * rate
    fee_inr = sale_inr * fee_pct / 100
    return dict(ship=ship, slab=slab, base=base, price=price, final=final, sale_inr=sale_inr,
                fee_inr=fee_inr, real_profit=sale_inr - fee_inr - cost - ship)


def price_calculator():
    st.subheader("💰 eBay Price Calculator")
    st.caption("Item cost + profit + shipping + eBay fee → listing price ending in .99. "
               "Shipping uses the ITPS rate table (GST included), same as the reference calculator.")
    mp = st.selectbox("Marketplace", list(CALC_MARKETS), key="calc_mp")
    m = CALC_MARKETS[mp]
    left, right = st.columns([1, 1], gap="large")
    with left:
        a, b = st.columns(2)
        cost = a.number_input("Item cost (INR)", min_value=0.0, value=450.0, step=10.0, key="calc_cost")
        profit = b.number_input("Desired profit (INR)", min_value=0.0, value=300.0, step=10.0, key="calc_profit")
        c, d = st.columns(2)
        weight = c.number_input("Shipping weight (grams)", min_value=0, value=250, step=50, key="calc_wt")
        fee = d.number_input("eBay fee (%)", min_value=0.0, max_value=99.0, value=15.0, step=0.5,
                             key=f"calc_fee_{mp}", help="Placeholder 15%. Set your real fee.")
        rate = st.number_input(f"Exchange rate (INR per {m['cur']})", min_value=0.01, value=m["rate"], step=1.0,
                               key=f"calc_rate_{mp}")
        ship_ov = st.number_input("Shipping override in INR (0 = use rate table)", min_value=0.0, value=0.0,
                                  step=10.0, key="calc_ship_ov")
    r = calc_listing_price(cost, profit, weight, fee, rate, m, ship_ov or None)
    with right:
        st.markdown("**Final listing price (.99)**")
        st.markdown(f"<div style='font-size:2.6rem;font-weight:700;color:#2b3ea0'>{m['sym']}{r['final']:,.2f} "
                    f"<span style='font-size:1rem;color:#888'>{m['cur']}</span></div>", unsafe_allow_html=True)
        rows = [
            ("Item cost", f"₹{cost:,.2f}"),
            (f"Shipping ({mp}" + (", override)" if ship_ov else f", up to {r['slab']} g)"), f"₹{r['ship']:,.2f}"),
            ("Desired profit", f"₹{profit:,.2f}"),
            ("Base cost (cost + profit + shipping)", f"₹{r['base']:,.2f}"),
            ("Price before .99 rounding", f"{m['sym']}{r['price']:,.2f}"),
            ("Sale value at final price", f"₹{r['sale_inr']:,.0f}"),
            (f"eBay fee ({fee:g}%)", f"- ₹{r['fee_inr']:,.0f}"),
            ("Profit at final price", f"₹{r['real_profit']:,.0f}"),
        ]
        st.table(pd.DataFrame(rows, columns=["", "INR / price"]).set_index(""))
        if r["real_profit"] >= profit - 0.5:
            st.success("Meets your profit target")
        else:
            st.warning("Below your profit target")
        st.caption("Assumes free shipping to the buyer, so shipping and the eBay fee come out of the listing price.")
    with st.expander("Shipping rates reference (INR, ITPS, GST included)"):
        wts = list(range(50, 5001, 50))
        tbl = pd.DataFrame({"Weight (g)": wts, **{
            name: [int(round_half_up(CALC_MARKETS[name]["s0"] + CALC_MARKETS[name]["ds"] * (w // 50 - 1), 0))
                   for w in wts] for name in ["USA", "UK", "Canada", "Australia"]}})
        st.dataframe(tbl, hide_index=True, use_container_width=True, height=380)


def main():
    st.set_page_config(page_title="Sharmex Global Tools", page_icon="📦", layout="wide")
    tab_labels, tab_calc = st.tabs(["📦 India Post Label Generator", "💰 eBay Price Calculator"])
    with tab_labels:
        label_generator()
    with tab_calc:
        price_calculator()


def label_generator():
    st.title("📦 India Post Bulk Label Generator")
    st.caption("Ship-To address always used · Multi-buy aware (reads Variation details) · Readable item names")

    with st.sidebar:
        st.header("Label generator settings")
        rates = {
            "GBP": st.number_input("GBP → INR", value=DEFAULT_RATES["GBP"], step=1.0),
            "USD": st.number_input("USD → INR", value=DEFAULT_RATES["USD"], step=1.0),
        }
        rates = {**DEFAULT_RATES, **rates}
        max_name = st.slider("Max item-name length", 25, 100, 50,
                             help="Length of HS DESCRIPTION. Lower it if India Post rejects long text.")
        lookup = st.checkbox("Look up missing county online", value=True,
                             help="Uses postcodes.io (UK) / zippopotam.us (US). Filled cells are sky blue - verify them.")
        ov_file = st.file_uploader("overrides.xlsx (optional)", type=["xlsx"],
                                   help="Your permanent corrections: item_id, pack, name, hs_code, weight_g, value_inr")
        hist_file = st.file_uploader("order_history.csv (optional)", type=["csv"],
                                     help="History saved from earlier runs. Confirmed weights/prices are re-used automatically.")

    show_legend()

    up = st.file_uploader("Upload eBay Orders CSV", type=["csv"])
    if not up:
        return
    if not os.path.exists(TEMPLATE_PATH):
        st.error("template.xlsx not found next to app.py")
        return
    overrides = {}
    if ov_file is not None:
        overrides = load_overrides(ov_file)
    elif os.path.exists(OVERRIDES_PATH):
        overrides = load_overrides(OVERRIDES_PATH)

    hist = load_history(HISTORY_PATH) if os.path.exists(HISTORY_PATH) else load_history(None)
    if hist_file is not None:
        hist = merge_history(hist, load_history(hist_file))
    learned = learn_from_history(hist)

    try:
        df, cols = read_orders_csv(up.getvalue())
        if df.empty:
            st.warning("No valid eBay orders found in the file.")
            return
        cfg = {"rates": rates, "max_name": max_name, "lookup": lookup}
        parcels0, sub_rows0, sub_marks0, marks, lines_out = process(df, cols, cfg, overrides, learned)
        _, lines0 = build_report(parcels0, lines_out)
    except Exception as e:  # noqa
        st.error(f"Error processing orders: {e}")
        raise

    # ---- review & adjust items (edits are applied to the Excel and can be saved as history)
    n_est = int(lines0["Source"].isin(ESTIMATED_SOURCES).sum())
    n_learned = int(lines0["Source"].astype(str).str.startswith("Learned").sum())
    with st.expander(f"✏️ Review & adjust item weights / HS / prices  ·  {n_est} estimated, {n_learned} learned from history",
                     expanded=n_est > 0):
        st.caption("Edit the white columns if a weight, HS code, value or name is wrong, or tick Confirmed if it is "
                   "right. Edited/confirmed items are saved to history and used automatically next time.")
        ed = lines0[["Serial", "Order", "Item number", "eBay title", "Qty", "Pack", "Source", "India Post name", "HS",
                     "Unit wt (g)", "Unit INR"]].copy()
        ed["Confirmed"] = lines0["Source"].astype(str).str.startswith(("Learned", "Override"))
        key = f"lines_editor_{getattr(up, 'name', '')}_{getattr(up, 'size', len(lines0))}_{len(lines0)}"
        edited = st.data_editor(
            ed, key=key, hide_index=True, use_container_width=True,
            disabled=["Serial", "Order", "Item number", "eBay title", "Qty", "Pack", "Source"],
            column_config={
                "Unit wt (g)": st.column_config.NumberColumn(min_value=0, step=1, help="Weight of ONE line unit incl. pack"),
                "Unit INR": st.column_config.NumberColumn(min_value=0, step=1, help="Declared INR value of one line unit"),
                "Confirmed": st.column_config.CheckboxColumn(help="Tick = these values are correct, learn them"),
            })
    parcels, sub_rows, sub_marks, final_lines = apply_edits(parcels0, sub_rows0, lines0, edited)
    out_bytes, added_cols = build_workbook(parcels, sub_rows, sub_marks, marks)
    orders_df, _ = build_report(parcels, lines_out)
    lines_df = final_lines
    if added_cols:
        st.warning("Your template.xlsx has no column for: " + ", ".join(added_cols) +
                   ". They were added at the end. If India Post rejects the file, add these headers to template.xlsx "
                   "in the right place.")

    n_chk = int((orders_df["Status"] == "CHECK").sum())
    n_rev = int((orders_df["Status"] == "Review").sum())
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Parcels", len(orders_df))
    c2.metric("Line items", len(lines_df))
    c3.metric("🟥 Must fix", n_chk)
    c4.metric("🟨 Review", n_rev)

    ts = pd.Timestamp.now().strftime("%Y%m%d_%H%M%S")
    d1, d2 = st.columns(2)
    d1.download_button("📥 Download India Post Excel", out_bytes, f"IndiaPost_BulkUpload_{ts}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
    rep = io.BytesIO()
    with pd.ExcelWriter(rep, engine="openpyxl") as w:
        orders_df.drop(columns="_order_key").to_excel(w, sheet_name="Parcels", index=False)
        lines_df.to_excel(w, sheet_name="Line items", index=False)
        ot = overrides_template(lines_df)
        if not ot.empty:
            ot.to_excel(w, sheet_name="overrides_to_fill", index=False)
    d2.download_button("📄 Download Review Report", rep.getvalue(), f"Review_{ts}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # ---- history
    merged = merge_history(hist, history_from_lines(lines_df))
    h1, h2 = st.columns(2)
    if h1.button("💾 Save this run to history (so the tool learns)"):
        try:
            merged.to_csv(HISTORY_PATH, index=False)
            msg = f"Saved {len(merged)} history rows to {HISTORY_PATH}."
            if github_enabled():
                github_push(HISTORY_PATH, merged.to_csv(index=False).encode(), f"Update order history {ts}")
                msg += " Also committed to GitHub."
            else:
                msg += " (On Streamlit Cloud this file is lost when the app restarts. Download it below and upload it to "
                msg += "GitHub, or set GITHUB_TOKEN/GITHUB_REPO in Streamlit secrets for automatic saving.)"
            st.success(msg)
        except Exception as e:  # noqa
            st.error(f"Could not save history: {e}")
    h2.download_button("⬇️ Download history file", merged.to_csv(index=False).encode(), "order_history.csv",
                       mime="text/csv")
    st.caption(f"History: {len(hist)} saved rows · {len(learned['exact'])} confirmed items learned")

    t1, t2, t3 = st.tabs(["⚠️ Needs attention", "All parcels", "Line items (what each parcel contains)"])
    show = orders_df.drop(columns="_order_key")
    with t1:
        att = show[show["Status"] != "OK"].copy()
        att = att.iloc[att["Status"].map({"CHECK": 0, "Review": 1}).argsort(kind="stable")]
        if att.empty:
            st.success("Nothing needs attention 🎉")
        else:
            st.dataframe(att[["Serial", "Order", "Status", "Ship-To name", "Address 1", "City", "Postcode",
                              "Address source", "Notes"]], use_container_width=True, hide_index=True)
    with t2:
        st.dataframe(show.drop(columns=["Notes"]), use_container_width=True, hide_index=True)
    with t3:
        st.dataframe(lines_df, use_container_width=True, hide_index=True)

if __name__ == "__main__":
    main()
